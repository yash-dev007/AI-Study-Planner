"""Safe extraction of text from uploaded syllabus documents.

The extractor never executes, renders, or otherwise trusts document content.
It only converts bytes to plain text and returns metadata. AI extracts structure
from the resulting text.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from app.core.config import settings
from app.core.errors import DocumentProcessingError

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}
CONTENT_TYPE_EXTENSIONS = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "text/plain": ".txt",
    "application/octet-stream": None,
}
_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")

# Byte-order marks for text encodings that legitimately contain NUL bytes.
# UTF-32 entries must be tested first: their BOM starts with the UTF-16 one.
_TEXT_BOMS = (
    (b"\xff\xfe\x00\x00", "utf-32"),
    (b"\x00\x00\xfe\xff", "utf-32"),
    (b"\xff\xfe", "utf-16"),
    (b"\xfe\xff", "utf-16"),
)
_TEXT_BOM_BYTES = tuple(bom for bom, _encoding in _TEXT_BOMS)
_EXCESS_BLANK_LINES = re.compile(r"\n{3,}")

# A DOCX is a zip archive: a few hundred kilobytes can inflate to gigabytes.
# These limits are checked from the archive directory before anything is inflated.
MAX_DOCX_ENTRIES = 2_000
MAX_DOCX_UNCOMPRESSED_BYTES = 60 * 1024 * 1024
MAX_DOCX_COMPRESSION_RATIO = 200
# Bounds the per-page text extraction work for very long PDFs.
MAX_PDF_PAGES = 400


@dataclass
class ExtractionResult:
    text: str
    page_count: int = 0
    warnings: list[str] = field(default_factory=list)


def detect_extension(filename: str, content_type: str) -> str:
    suffix = Path(filename or "").suffix.lower()
    if suffix in SUPPORTED_EXTENSIONS:
        return suffix
    mapped = CONTENT_TYPE_EXTENSIONS.get((content_type or "").split(";")[0].strip().lower())
    if mapped:
        return mapped
    raise DocumentProcessingError(
        "Unsupported file type. Upload a PDF, DOCX, or TXT syllabus.",
        code="UNSUPPORTED_FILE_TYPE",
        details={"filename": filename, "content_type": content_type},
    )


def validate_upload(filename: str, content_type: str, data: bytes) -> str:
    max_bytes = settings.max_upload_mb * 1024 * 1024
    if not data:
        raise DocumentProcessingError("The uploaded file is empty.", code="EMPTY_FILE")
    if len(data) > max_bytes:
        raise DocumentProcessingError(
            f"The file exceeds the {settings.max_upload_mb} MB upload limit.",
            code="FILE_TOO_LARGE",
            details={"size_bytes": len(data), "limit_bytes": max_bytes},
        )

    extension = detect_extension(filename, content_type)
    if extension == ".pdf" and not data.startswith(b"%PDF"):
        raise DocumentProcessingError(
            "The file extension is .pdf but the content is not a valid PDF.",
            code="CORRUPT_DOCUMENT",
        )
    if extension == ".docx" and not data.startswith(b"PK"):
        raise DocumentProcessingError(
            "The file extension is .docx but the content is not a valid DOCX archive.",
            code="CORRUPT_DOCUMENT",
        )
    if extension == ".txt":
        sample = data[:2048]
        # NUL bytes mean binary content, except in text encoded as UTF-16/UTF-32,
        # which starts with a byte-order mark.
        if b"\x00" in sample and not sample.startswith(_TEXT_BOM_BYTES):
            raise DocumentProcessingError(
                "The text file appears to be binary and cannot be read as a syllabus.",
                code="CORRUPT_DOCUMENT",
            )
    return extension


def _check_docx_archive(data: bytes) -> None:
    """Reject a DOCX archive that would inflate past the memory budget.

    Checked from the archive directory (declared sizes) so a zip bomb is refused
    before python-docx inflates it.
    """

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_DOCX_ENTRIES:
                raise DocumentProcessingError(
                    "The DOCX file contains an implausible number of parts and was rejected.",
                    code="CORRUPT_DOCUMENT",
                    details={"entries": len(infos)},
                )
            total = 0
            for info in infos:
                total += info.file_size
                if total > MAX_DOCX_UNCOMPRESSED_BYTES:
                    raise DocumentProcessingError(
                        "The DOCX file expands beyond the size this server will process.",
                        code="CORRUPT_DOCUMENT",
                        details={"declared_bytes": total},
                    )
                if (
                    info.compress_size > 0
                    and info.file_size / info.compress_size > MAX_DOCX_COMPRESSION_RATIO
                ):
                    raise DocumentProcessingError(
                        "The DOCX file is compressed far beyond a normal document and was "
                        "rejected.",
                        code="CORRUPT_DOCUMENT",
                        details={"entry": info.filename},
                    )
    except zipfile.BadZipFile as exc:
        raise DocumentProcessingError(
            "The DOCX file could not be read. It may be corrupted or password protected.",
            code="CORRUPT_DOCUMENT",
        ) from exc


def sanitize_text(text: str) -> str:
    cleaned = text.replace("\ufeff", "")
    cleaned = cleaned.replace("\r\n", "\n").replace("\r", "\n")
    cleaned = _CONTROL_CHARS.sub(" ", cleaned)
    cleaned = _EXCESS_BLANK_LINES.sub("\n\n", cleaned)
    cleaned = "\n".join(line.rstrip() for line in cleaned.split("\n"))
    if len(cleaned) > settings.max_extracted_chars:
        cleaned = cleaned[: settings.max_extracted_chars]
    return cleaned.strip()



def extract_text(filename: str, content_type: str, data: bytes) -> ExtractionResult:
    extension = validate_upload(filename, content_type, data)

    if extension == ".txt":
        text = _decode_text(data)
        return ExtractionResult(text=_require_text(sanitize_text(text), ".txt"))

    if extension == ".docx":
        text = _extract_docx(data)
        return ExtractionResult(text=_require_text(sanitize_text(text), ".docx"))

    return _extract_pdf(data)


def _require_text(text: str, extension: str) -> str:
    if not text.strip():
        raise DocumentProcessingError(
            f"The {extension} file contains no readable text. Check the file and try again.",
            code="NO_TEXT_EXTRACTED",
        )
    return text


def _decode_text(data: bytes) -> str:
    # Honour the byte-order mark first: decoding UTF-32 as UTF-16 silently yields
    # NUL-riddled noise instead of the document text.
    for bom, encoding in _TEXT_BOMS:
        if data.startswith(bom):
            return data.decode(encoding, errors="replace")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _extract_docx(data: bytes) -> str:
    _check_docx_archive(data)
    try:
        import docx

        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentProcessingError(
            "The DOCX file could not be read. It may be corrupted or password protected.",
            code="CORRUPT_DOCUMENT",
            details={"reason": str(exc)},
        ) from exc

    lines: list[str] = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(line for line in lines if line is not None)


def _extract_pdf(data: bytes) -> ExtractionResult:
    try:
        import fitz

        document = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise DocumentProcessingError(
            "The PDF could not be opened. It may be corrupted or password protected.",
            code="CORRUPT_DOCUMENT",
            details={"reason": str(exc)},
        ) from exc

    with document:
        if document.needs_pass:
            raise DocumentProcessingError(
                "The PDF is password protected. Upload an unprotected copy.",
                code="PASSWORD_PROTECTED_DOCUMENT",
            )
        page_count = document.page_count
        if page_count == 0:
            raise DocumentProcessingError("The PDF contains no pages.", code="CORRUPT_DOCUMENT")

        pages: list[str] = []
        warnings: list[str] = []
        characters = 0
        for number in range(page_count):
            if number >= MAX_PDF_PAGES:
                warnings.append(
                    f"Only the first {MAX_PDF_PAGES} of {page_count} pages were read."
                )
                break
            text = document.load_page(number).get_text("text")
            pages.append(text)
            characters += len(text)
            if characters >= settings.max_extracted_chars:
                # Stop paying for pages that would be truncated anyway.
                if number + 1 < page_count:
                    warnings.append(
                        "The document is longer than the extraction limit; later pages were "
                        "not read."
                    )
                break

        text = "\n\n".join(pages)
        if not text.strip():
            raise DocumentProcessingError(
                "No readable text could be extracted from the PDF. Upload a text-based "
                "document or paste your syllabus text directly.",
                code="NO_TEXT_EXTRACTED",
                details={"page_count": page_count},
            )
        return ExtractionResult(
            text=sanitize_text(text),
            page_count=page_count,
            warnings=warnings,
        )
