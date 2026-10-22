"""Document ingestion tests: validation, PDF/DOCX/TXT extraction, and corrupt input."""

from __future__ import annotations

import io

import pytest
from app.core.errors import DocumentProcessingError
from app.documents.extractor import extract_text, sanitize_text, validate_upload

SAMPLE = """\
Unit 1: Introduccion a los sistemas operativos
1. Definicion y funciones (60 minutes)
2. Historia y evolucion (45 minutes)

Unit 2: Procesos
- Procesos y estados
- Planificacion de CPU
"""


def make_pdf(pages: list[str]) -> bytes:
    """A real PDF, built in-memory: text pages have embedded text, empty ones do not."""

    import fitz

    document = fitz.open()
    for text in pages:
        page = document.new_page()
        if text:
            page.insert_text((72, 96), text, fontsize=11)
    data = document.tobytes()
    document.close()
    return data


def make_docx(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table_rows:
        table = document.add_table(rows=0, cols=len(table_rows[0]))
        for row in table_rows:
            cells = table.add_row().cells
            for index, value in enumerate(row):
                cells[index].text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
# Upload validation
# --------------------------------------------------------------------------- #


def test_valid_upload_returns_extension() -> None:
    assert validate_upload("syllabus.txt", "text/plain", b"hello") == ".txt"
    assert validate_upload("Syllabus.PDF", "application/pdf", b"%PDF-1.7 x") == ".pdf"
    assert (
        validate_upload(
            "syllabus.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            b"PK\x03\x04rest",
        )
        == ".docx"
    )


def test_unsupported_extension_is_rejected() -> None:
    with pytest.raises(DocumentProcessingError) as excinfo:
        validate_upload("plan.xlsx", "application/vnd.ms-excel", b"PK\x03\x04")
    assert excinfo.value.code == "UNSUPPORTED_FILE_TYPE"


def test_empty_file_is_rejected() -> None:
    with pytest.raises(DocumentProcessingError) as excinfo:
        validate_upload("syllabus.txt", "text/plain", b"")
    assert excinfo.value.code == "EMPTY_FILE"


def test_file_that_lies_about_its_type_is_rejected() -> None:
    with pytest.raises(DocumentProcessingError) as excinfo:
        validate_upload("syllabus.pdf", "application/pdf", b"this is plain text")
    assert excinfo.value.code == "CORRUPT_DOCUMENT"


def test_empty_text_file_is_rejected_with_clear_code() -> None:
    with pytest.raises(DocumentProcessingError) as excinfo:
        extract_text("syllabus.txt", "text/plain", b"   \n \t \n")
    assert excinfo.value.code == "NO_TEXT_EXTRACTED"


def test_sanitize_text_collapses_noise() -> None:
    cleaned = sanitize_text("  Unit 1  \x00\n\n\n\n\nTopic A\r\nTopic B   \n")
    assert "\x00" not in cleaned
    assert "\n\n\n\n" not in cleaned
    assert cleaned.endswith("Topic B")


# --------------------------------------------------------------------------- #
# TXT
# --------------------------------------------------------------------------- #


def test_text_file_is_decoded() -> None:
    result = extract_text("syllabus.txt", "text/plain", SAMPLE.encode("utf-8"))
    assert result.page_count == 0, "a text file has no pages"
    assert "Unit 1" in result.text


def test_utf16_text_file_is_decoded() -> None:
    result = extract_text("syllabus.txt", "text/plain", SAMPLE.encode("utf-16"))
    assert "Planificacion de CPU" in result.text


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #


def test_docx_paragraphs_and_tables_are_extracted() -> None:
    data = make_docx(
        ["Unit 1: Memory", "1. Paging and segmentation (90 minutes)"],
        table_rows=[["Topic", "Duration"], ["Virtual memory", "120 minutes"]],
    )
    result = extract_text(
        "syllabus.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        data,
    )
    assert "Unit 1: Memory" in result.text
    assert "Topic | Duration" in result.text
    assert "Virtual memory | 120 minutes" in result.text


def test_corrupted_docx_is_reported() -> None:
    with pytest.raises(DocumentProcessingError) as excinfo:
        extract_text(
            "syllabus.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            b"PK\x03\x04not-really-a-docx",
        )
    assert excinfo.value.code == "CORRUPT_DOCUMENT"


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #


def test_pdf_text_layer_is_used() -> None:
    result = extract_text("syllabus.pdf", "application/pdf", make_pdf([SAMPLE]))
    assert result.page_count == 1
    assert "Unit 1" in result.text


def test_image_only_pdf_fails_with_actionable_error() -> None:
    """A scanned or image-only PDF with no embedded text must fail with a clear actionable error."""

    with pytest.raises(DocumentProcessingError) as excinfo:
        extract_text("scanned.pdf", "application/pdf", make_pdf(["", ""]))
    assert excinfo.value.code == "NO_TEXT_EXTRACTED"
    assert "no readable text" in excinfo.value.message.lower()


def test_mixed_pdf_extracts_readable_pages() -> None:
    """Readable pages are extracted even if some pages have no text."""

    result = extract_text("mixed.pdf", "application/pdf", make_pdf([SAMPLE, ""]))
    assert "Unit 1" in result.text
    assert result.page_count == 2


def test_blank_pdf_is_rejected() -> None:
    with pytest.raises(DocumentProcessingError) as excinfo:
        extract_text("blank.pdf", "application/pdf", make_pdf([""]))
    assert excinfo.value.code == "NO_TEXT_EXTRACTED"
